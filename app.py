"""
ISU Civil Engineering Lab Data Analyzer
Idaho State University — Civil Engineering Department
Version 1.0.0

Processes CSV exports from Campbell Scientific CR6 datalogger (LoggerNet).
Runs completely offline. Deployable as a standalone .exe via PyInstaller.

Usage:
    python app.py

Packaging:
    pyinstaller --onefile --windowed app.py
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import os
import sys
import io
import csv
import threading
import time
import glob
import shutil
import json
from datetime import datetime
from pathlib import Path

import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('TkAgg')

import matplotlib.dates as mdates
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg, NavigationToolbar2Tk
from matplotlib.figure import Figure


# ──────────────────────────────────────────────────────────────────────────────
# Constants
# ──────────────────────────────────────────────────────────────────────────────

APP_TITLE  = "ISU Civil Engineering Lab Data Analyzer"
VERSION    = "1.0.0"
RECENT_MAX = 8

def _app_dir() -> Path:
    """Return the folder where user data files (configs, recent) live.
    In a PyInstaller .exe, use the folder next to the .exe rather than
    the temp extraction directory."""
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).parent
    return Path(__file__).parent

CONFIGS_FILE = _app_dir() / 'graph_configs.json'
RECENT_FILE  = _app_dir() / 'recent_files.json'

# Written by camera_trigger.py — motion events + live camera status.
# See camera_trigger.py for the file schema.
TRIGGER_FILE      = _app_dir() / 'trigger_event.json'
CAM_PREVIEW_FILES = {1: _app_dir() / 'camera_preview_1.png',
                     2: _app_dir() / 'camera_preview_2.png'}

# Light/professional color palette (ISU colors — navy + white)
C = {
    'bg':         '#F0F2F5',
    'panel':      '#FFFFFF',
    'border':     '#D1D5DB',
    'accent':     '#1E3A5F',
    'accent_lt':  '#EBF2FA',
    'btn_primary':'#2563EB',
    'btn_danger': '#DC2626',
    'btn_normal': '#E5E7EB',
    'text':       '#111827',
    'muted':      '#6B7280',
    'success':    '#16A34A',
}

# Colors cycled when plotting multiple series
PLOT_PALETTE = ['#2563EB', '#DC2626', '#16A34A', '#D97706', '#7C3AED', '#0891B2']

# Maps varied CSV column headers → normalized internal names.
# First alias that matches wins; comparison is case-insensitive.
COL_ALIASES = {
    'time_s':     ['time (s)', 'time(s)', 'time_s', 'time', 'timestamp', 'elapsed_s'],
    'load_lbf':   ['load (lbf)', 'load(lbf)', 'load_lbf', 'load'],
    'stress_psi': ['stress (psi)', 'stress(psi)', 'stress_psi', 'stress', 'ch2'],
    'ch3':        ['ch3', 'displacement (in)', 'displacement'],
    'ch4':        ['ch4', 'strain'],
    # Vibrating-Wire Live Monitor CSV columns (CellA_Load_lbf, etc.)
    'cella_load_lbf':  ['cella_load_lbf', 'cell a load (lbf)', 'cell_a_load'],
    'cellb_load_lbf':  ['cellb_load_lbf', 'cell b load (lbf)', 'cell_b_load'],
    'cella_r1_digits': ['cella_r1_digits', 'cell a r1 digits', 'cell_a_r1'],
    'cellb_r1_digits': ['cellb_r1_digits', 'cell b r1 digits', 'cell_b_r1'],
}

# Human-readable labels for normalized column names
COL_LABELS = {
    'time_s':     'Time (s)',
    'load_lbf':   'Load (lbf)',
    'stress_psi': 'Stress (psi)',
    'ch3':        'CH3',
    'ch4':        'CH4',
    'cella_load_lbf':  'Cell A Load (lbf)',
    'cellb_load_lbf':  'Cell B Load (lbf)',
    'cella_r1_digits': 'Cell A R1 (digits)',
    'cellb_r1_digits': 'Cell B R1 (digits)',
    'elapsed_s':       'Elapsed Time (s)',
}

# Graph presets: (display name, x_col, y_col, x_label, y_label)
# y_col == '__all__' plots every numeric column on one chart
PRESETS = [
    ("— Select Preset —",   None,       None,         '',            ''),
    ("Load vs Time",        'time_s',   'load_lbf',   'Time (s)',    'Load (lbf)'),
    ("Stress vs Time",      'time_s',   'stress_psi', 'Time (s)',    'Stress (psi)'),
    ("Load vs Displacement",'ch3',      'load_lbf',   'Disp. (in)', 'Load (lbf)'),
    ("Stress vs Strain",    'ch4',      'stress_psi', 'Strain',      'Stress (psi)'),
    ("Load vs CH4",         'ch4',      'load_lbf',   'CH4',         'Load (lbf)'),
    ("All Channels vs Time",'time_s',   '__all__',    'Time (s)',    'Value'),
    ("VW — Cell A Load vs Time",  'time_s',         'cella_load_lbf',
     'Time (s)',            'Cell A Load (lbf)'),
    ("VW — Cell B Load vs Time",  'time_s',         'cellb_load_lbf',
     'Time (s)',            'Cell B Load (lbf)'),
    ("VW — Both Cells vs Time",   'time_s',         '__all__',
     'Time (s)',            'Load (lbf)'),
    ("VW — Cell A vs Cell B",     'cella_load_lbf', 'cellb_load_lbf',
     'Cell A Load (lbf)',   'Cell B Load (lbf)'),
]

CHART_TYPES = ['Line', 'Scatter', 'Area']

ALL_CHANNELS_LABEL = '(all channels)'


# ──────────────────────────────────────────────────────────────────────────────
# Utility
# ──────────────────────────────────────────────────────────────────────────────

def normalize_df(df: pd.DataFrame, time_origin=None) -> pd.DataFrame:
    """
    Normalize a raw DataFrame into the app's standard column schema.

    1. Alias known column names (case-insensitive, strip-safe) — best-effort.
    2. Convert any non-numeric time column (datetime strings or objects) to
       elapsed seconds from the first valid row — or from ``time_origin``
       (an absolute timestamp) when given, so successive live files share
       one time axis instead of each restarting at 0.
    3. Coerce every column to numeric (non-parseable values become NaN).
    4. Drop columns that are entirely NaN after coercion.
    5. Drop rows where time_s is NaN (if time_s is present).
    Returns whatever numeric columns survive — never raises.  When step 2
    converted timestamps, the origin used is stored in
    ``df.attrs['time_origin']`` so callers can align later files to it.
    """
    if df is None or df.empty:
        return pd.DataFrame()

    # Step 1 — build case-insensitive lookup and rename to standard names
    lower_map = {c.lower().strip(): c for c in df.columns}
    rename = {}
    for std_name, aliases in COL_ALIASES.items():
        for alias in aliases:
            key = alias.strip().lower()
            if key in lower_map and lower_map[key] not in rename:
                rename[lower_map[key]] = std_name
                break
    df = df.rename(columns=rename)

    # Step 2 — convert time column to elapsed seconds
    if 'time_s' in df.columns:
        if not pd.api.types.is_numeric_dtype(df['time_s']):
            # Datetime strings (e.g. "2026-08-27 23:55:30"). format='mixed'
            # parses per-element: TOA5 trims zero fractional seconds, so
            # sub-second tables mix "…:20.25" with "…:21" and inferring one
            # format from the first row silently drops the other kind.
            parsed = pd.to_datetime(df['time_s'], format='mixed',
                                    errors='coerce')
            if parsed.notna().any():
                origin = (time_origin if time_origin is not None
                          else parsed.dropna().iloc[0])
                df['time_s'] = (parsed - origin).dt.total_seconds()
                df.attrs['time_origin'] = origin
        else:
            # Excel serial-date numbers produced by the pyxlsb engine:
            # values like 46261.99... (days since 1899-12-30) are >> 40000.
            # Regular elapsed-seconds values from a lab test are never that large.
            vals = df['time_s'].dropna()
            if len(vals) > 0 and float(vals.mean()) > 40000:
                as_dt = pd.to_datetime(df['time_s'], unit='D',
                                       origin='1899-12-30', errors='coerce')
                if as_dt.notna().any():
                    origin = as_dt.dropna().iloc[0]
                    df['time_s'] = (as_dt - origin).dt.total_seconds()

    # Step 3 — coerce every column to numeric
    for col in list(df.columns):
        df[col] = pd.to_numeric(df[col], errors='coerce')

    # Step 4 — drop columns that are entirely NaN (non-numeric metadata rows)
    df = df.dropna(axis=1, how='all')

    # Step 5 — drop rows where the time axis is NaN
    if 'time_s' in df.columns:
        df = df.dropna(subset=['time_s'])

    return df.reset_index(drop=True)


def detect_format(path) -> str:
    """Return a human-readable label for a data file based on its extension."""
    ext = Path(path).suffix.lower()
    if ext == '.xlsx':
        return 'Excel (.xlsx)'
    if ext == '.dat':
        return 'TOA5 (.dat)'
    return 'CSV'


def _read_csv_robust(path) -> pd.DataFrame:
    """
    Try multiple CSV parsing strategies and return the one that produces the
    most numeric data.  Strategies: comma then auto-detect delimiter, each
    tried with 0-4 skipped header rows (handles TOA5-format CSVs, metadata
    preambles, unit/processing rows, etc.).  Falls back to a headerless parse
    if nothing else works.
    """
    best_df = None
    best_score = -1

    for sep, engine in ((',', 'c'), (None, 'python')):
        for skip in range(5):
            try:
                kw = dict(skiprows=skip, skip_blank_lines=True)
                if sep is None:
                    kw['sep'] = None
                    kw['engine'] = 'python'
                df = pd.read_csv(path, **kw)
                if df.empty or len(df.columns) < 1:
                    continue
                # Score: coercible numeric columns × row count
                tmp = df.copy()
                for c in tmp.columns:
                    tmp[c] = pd.to_numeric(tmp[c], errors='coerce')
                n_numeric = tmp.dropna(axis=1, how='all').shape[1]
                score = n_numeric * len(df)
                if score > best_score:
                    best_score = score
                    best_df = df
            except Exception:
                continue

    if best_df is not None:
        return best_df

    # Last resort: treat every row as data with no header
    try:
        return pd.read_csv(path, header=None, skip_blank_lines=True)
    except Exception as e:
        raise ValueError(f"Could not parse as CSV: {e}")


def read_file(path, time_origin=None) -> pd.DataFrame:
    """
    Read a data file into a normalized DataFrame.

    Dispatch:
      .xlsx  → pd.read_excel (openpyxl)
      .dat   → Campbell Scientific TOA5 parser; falls back to CSV if not TOA5
      other  → CSV with automatic delimiter detection

    Every path calls normalize_df() for consistent aliasing, coercion, and
    time handling.  Never raises on a partially-readable file — returns
    whatever numeric columns survive normalization.
    """
    ext = Path(path).suffix.lower()

    if ext == '.xlsx':
        raw = pd.read_excel(path, engine='openpyxl')

    elif ext == '.xlsb':
        raw = pd.read_excel(path, engine='pyxlsb')

    elif ext == '.dat':
        with open(path, 'r', newline='', errors='ignore') as f:
            text = f.read()
        raw = _read_toa5(text)
        if raw is None or raw.empty:
            raw = _read_csv_robust(path)

    else:
        raw = _read_csv_robust(path)

    return normalize_df(raw, time_origin=time_origin)


# ──────────────────────────────────────────────────────────────────────────────
# Vibrating-Wire load-cell support (used by VWLiveMonitor)
# ──────────────────────────────────────────────────────────────────────────────

# Default calibration constants for a vibrating-wire load cell.
# Linear:     Load = G * (R1 - R0)
# Polynomial: Load = A * R1^2 + B * R1 + C
VW_DEFAULTS = {
    'R0': 7319.0,
    'G':  -96.20,
    'A':  -0.001148,
    'B':  -82.04,
    'C':  660800.0,
}

# Default gauge frequency column names in the CR6 TOA5 table.
VW_FREQ_COLS = ['Freq_1', 'Freq_2', 'Freq_3']


def _read_toa5(text: str) -> pd.DataFrame:
    """
    Parse a Campbell Scientific TOA5 table (as text) into a DataFrame.

    TOA5 layout:
        line 0  — environment line ("TOA5", logger info …)
        line 1  — field names      ("TIMESTAMP","RECORD","Freq_1", …)
        line 2  — units            ("TS","RN","Hz", …)
        line 3  — processing       ("","","Smp", …)
        line 4+ — data rows
    TIMESTAMP is parsed to datetime; rows with an unparseable timestamp are
    dropped (guards against a torn/partial row slipping through). Other columns
    are left as read so callers coerce only the columns they use.
    """
    lines = text.splitlines()
    if len(lines) < 5:
        return pd.DataFrame()

    # The header line is itself quoted/comma-separated — parse it with csv so
    # embedded commas/quotes are handled correctly.
    try:
        header = next(csv.reader([lines[1]]))
    except Exception:
        return pd.DataFrame()

    body = "\n".join(lines[4:])
    if not body.strip():
        return pd.DataFrame(columns=header)

    try:
        df = pd.read_csv(io.StringIO(body), header=None, names=header)
    except Exception:
        return pd.DataFrame(columns=header)

    if 'TIMESTAMP' in df.columns:
        # format='mixed' (per-element): TOA5 trims zero fractional seconds,
        # so sub-second tables mix "…:20.25" with "…:21"; inferring one
        # format from the first row would NaT-out the other kind.
        df['TIMESTAMP'] = pd.to_datetime(df['TIMESTAMP'], format='mixed',
                                         errors='coerce')
        df = df.dropna(subset=['TIMESTAMP']).reset_index(drop=True)
    return df


def read_data_snapshot(path: str) -> pd.DataFrame:
    """
    Read a growing LoggerNet TOA5 .dat file without blocking LoggerNet.

    Opens the file read-only (which does not take a write lock, so LoggerNet
    can keep appending), reads its current contents, and trims a partial last
    line — LoggerNet may be mid-write when we read, leaving the final line
    incomplete. Returns the parsed DataFrame (possibly empty).
    """
    with open(path, 'r', newline='', errors='ignore') as f:
        data = f.read()

    # Drop a trailing partial line (no terminating newline => still being written).
    if not data.endswith('\n'):
        idx = data.rfind('\n')
        data = data[:idx + 1] if idx != -1 else ''

    return _read_toa5(data)


def vw_compute(df: pd.DataFrame, g_cols: list, input_is_digits: bool,
               equation: str, consts: dict):
    """
    Compute (R1, Load) for every row of a TOA5 DataFrame.

    g_cols          — the three gauge frequency columns to average.
    input_is_digits — True if the columns already hold digits; otherwise the
                      values are Hz and are converted: digits = Hz^2 / 1000.
    R1              — per-row mean of the three gauge digit values.
    equation        — 'Linear' => Load = G*(R1 - R0)
                      'Polynomial' => Load = A*R1^2 + B*R1 + C
    consts          — dict of calibration constants (missing keys fall back to
                      VW_DEFAULTS).

    Returns (R1, Load) as pandas Series aligned to df.index. This is the single
    source of truth feeding the graph, the readout, AND the CSV log.
    """
    if df is None or len(df) == 0:
        empty = pd.Series([], dtype=float)
        return empty, empty

    digit_series = []
    for col in g_cols:
        s = pd.to_numeric(df[col], errors='coerce')
        if not input_is_digits:
            s = (s ** 2) / 1000.0
        digit_series.append(s)

    R1 = pd.concat(digit_series, axis=1).mean(axis=1)

    if equation == 'Polynomial':
        A = consts.get('A', VW_DEFAULTS['A'])
        B = consts.get('B', VW_DEFAULTS['B'])
        Cc = consts.get('C', VW_DEFAULTS['C'])
        Load = A * R1 ** 2 + B * R1 + Cc
    else:  # Linear
        R0 = consts.get('R0', VW_DEFAULTS['R0'])
        G = consts.get('G', VW_DEFAULTS['G'])
        Load = G * (R1 - R0)

    return R1, Load


# ──────────────────────────────────────────────────────────────────────────────
# Data Manager
# ──────────────────────────────────────────────────────────────────────────────

class DataManager:
    """
    Owns all data state. Handles CSV loading, column normalization,
    time chopping, backup, and export.

    df_full    — original merged data; never modified after a load
    df_current — working copy; may be chopped
    """

    def __init__(self):
        self.df_full        = None
        self.df_current     = None
        self.time_origin    = None   # absolute t=0 shared by all loaded files
        self.source_files   = []
        self.backup_dir     = Path.home() / 'Desktop' / 'LabData_Backup'
        self.compare_frames = []   # populated by load_separate()
        self.compare_names  = []   # short filename labels for legend

    # ── Loading ──────────────────────────────────────────────────────────────

    def load(self, paths: list) -> list:
        """
        Load and merge one or more CSV files.
        Returns a list of non-fatal error strings (empty = all good).
        Raises ValueError if nothing could be loaded at all.
        """
        frames, errors = [], []
        origin = None
        for p in paths:
            try:
                df = read_file(p, time_origin=origin)
                if origin is None:
                    origin = df.attrs.get('time_origin')
                frames.append(df)
            except Exception as e:
                errors.append(f"{Path(p).name}: {e}")

        if not frames:
            raise ValueError("Could not load any files.\n" + "\n".join(errors))

        merged = pd.concat(frames, ignore_index=True)
        if 'time_s' in merged.columns:
            merged = merged.sort_values('time_s').reset_index(drop=True)

        self.df_full      = merged.copy()
        self.df_current   = merged.copy()
        self.time_origin  = origin
        self.source_files = list(paths)
        return errors

    def append(self, paths: list) -> list:
        """
        Append new files to the existing dataset (for live mode).
        Falls back to a fresh load if nothing is loaded yet.
        """
        if self.df_full is None:
            return self.load(paths)

        frames, errors = [self.df_full], []
        origin = self.time_origin
        for p in paths:
            try:
                df = read_file(p, time_origin=origin)
                if origin is None:
                    origin = df.attrs.get('time_origin')
                frames.append(df)
            except Exception as e:
                errors.append(f"{Path(p).name}: {e}")

        merged = pd.concat(frames, ignore_index=True)
        if 'time_s' in merged.columns:
            merged = merged.sort_values('time_s').reset_index(drop=True)

        self.df_full      = merged.copy()
        self.df_current   = merged.copy()
        self.time_origin  = origin
        self.source_files.extend(paths)
        return errors

    def load_separate(self, paths: list) -> list:
        """
        Load files as independent series for side-by-side comparison.
        Each file keeps its own DataFrame instead of being merged.
        Also sets df_full/df_current to the first file so stats/chop still work.
        """
        frames, names, errors = [], [], []
        for p in paths:
            try:
                df = read_file(p)
                frames.append(df)
                names.append(Path(p).stem[:24])
            except Exception as e:
                errors.append(f"{Path(p).name}: {e}")

        if not frames:
            raise ValueError("Could not load any files.\n" + "\n".join(errors))

        self.compare_frames = frames
        self.compare_names  = names
        self.source_files   = list(paths)

        # Keep df_full/df_current pointing at frame[0] so backup/stats still work
        self.df_full    = frames[0].copy()
        self.df_current = frames[0].copy()
        return errors

    @property
    def compare_columns(self) -> list:
        """Columns present in ALL comparison frames (safe to plot across files)."""
        if not self.compare_frames:
            return self.columns
        sets = [set(df.columns) for df in self.compare_frames]
        return sorted(set.intersection(*sets))

    # ── Chopping ─────────────────────────────────────────────────────────────

    def chop(self, t_start: float, t_end: float) -> int:
        """Filter df_current to [t_start, t_end]. Returns retained row count."""
        if self.df_full is None:
            raise ValueError("No data loaded.")
        if 'time_s' not in self.df_full.columns:
            raise ValueError("No time column found in data.")

        mask   = (self.df_full['time_s'] >= t_start) & (self.df_full['time_s'] <= t_end)
        result = self.df_full[mask].copy().reset_index(drop=True)

        if len(result) == 0:
            raise ValueError(
                f"No data points found between {t_start}s and {t_end}s.\n"
                f"Full range: {self.time_range[0]:.2f}s – {self.time_range[1]:.2f}s"
            )

        self.df_current = result
        return len(result)

    def reset(self):
        """Restore df_current to the full unmodified dataset."""
        if self.df_full is None:
            raise ValueError("No data loaded.")
        self.df_current = self.df_full.copy()

    # ── Backup / Export ───────────────────────────────────────────────────────

    def backup(self, tag: str = "") -> str:
        """
        Save df_full (the full pre-chop dataset) to the backup folder.
        Returns the path to the written file.
        """
        if self.df_full is None:
            raise ValueError("No data to back up.")

        self.backup_dir.mkdir(parents=True, exist_ok=True)
        ts  = datetime.now().strftime('%Y%m%d_%H%M%S')
        sfx = f"_{tag}" if tag else ""
        out = self.backup_dir / f"backup_{ts}{sfx}.csv"
        self.df_full.to_csv(out, index=False)
        return str(out)

    def export(self, path: str):
        """Save df_current (post-chop working view) to the given path."""
        if self.df_current is None:
            raise ValueError("No data to export.")
        self.df_current.to_csv(path, index=False)

    # ── Statistics ────────────────────────────────────────────────────────────

    def stats(self) -> dict:
        """Return {col: {min, max, mean}} for all numeric columns in df_current."""
        if not self.has_data:
            return {}
        out = {}
        for col in self.df_current.select_dtypes(include=[np.number]).columns:
            s = self.df_current[col].dropna()
            if len(s):
                out[col] = {
                    'min':  round(float(s.min()),  3),
                    'max':  round(float(s.max()),  3),
                    'mean': round(float(s.mean()), 3),
                }
        return out

    # ── Properties ────────────────────────────────────────────────────────────

    @property
    def has_data(self) -> bool:
        return self.df_current is not None and len(self.df_current) > 0

    @property
    def columns(self) -> list:
        return list(self.df_current.columns) if self.df_current is not None else []

    @property
    def time_range(self) -> tuple:
        if not self.has_data or 'time_s' not in self.df_current.columns:
            return (0.0, 0.0)
        return (float(self.df_current['time_s'].min()),
                float(self.df_current['time_s'].max()))


# ──────────────────────────────────────────────────────────────────────────────
# Live Watcher
# ──────────────────────────────────────────────────────────────────────────────

class LiveWatcher:
    """
    Background thread that polls a folder every 1 second for new CSV files.
    Calls on_new_files(paths) when any appear. Thread-safe: only the callback
    touches shared state, and callers must route it through root.after().
    """

    def __init__(self, on_new_files):
        self._cb     = on_new_files
        self._stop   = threading.Event()
        self._thread = None
        self._known  = set()
        self.active  = False
        self.folder  = None

    def start(self, folder: str):
        if self.active:
            self.stop()
        self.folder = folder
        self._known = set(glob.glob(os.path.join(folder, '*.csv')))
        self._stop.clear()
        self.active  = True
        self._thread = threading.Thread(target=self._loop, daemon=True)
        self._thread.start()

    def stop(self):
        self._stop.set()
        self.active = False

    def _loop(self):
        while not self._stop.is_set():
            try:
                current = set(glob.glob(os.path.join(self.folder, '*.csv')))
                new     = sorted(current - self._known)
                if new:
                    self._known = current
                    self._cb(new)
            except Exception:
                pass
            time.sleep(1.0)


# ──────────────────────────────────────────────────────────────────────────────
# Trigger Watcher
# ──────────────────────────────────────────────────────────────────────────────

class TriggerWatcher:
    """
    Background thread that polls trigger_event.json (written by camera_trigger.py)
    once a second for new motion events. Each discrete motion start/stop bumps a
    `seq` counter in the file; when it rises, on_trigger(event) is called exactly
    once with the parsed event dict. Modeled after LiveWatcher — thread-safe:
    only the callback touches shared state, and callers must route it through
    root.after().
    """

    def __init__(self, on_trigger):
        self._cb       = on_trigger
        self._stop     = threading.Event()
        self._thread   = None
        self._last_seq = -1
        self.active    = False
        self.path      = None

    def start(self, path: str):
        if self.active:
            self.stop()
        self.path = path
        # Adopt the current seq so we don't fire on a stale pre-existing event.
        self._last_seq = self._read_seq()
        self._stop.clear()
        self.active  = True
        self._thread = threading.Thread(target=self._loop, daemon=True)
        self._thread.start()

    def stop(self):
        self._stop.set()
        self.active = False

    def _read_seq(self) -> int:
        try:
            data = json.loads(Path(self.path).read_text(encoding='utf-8'))
            return int(data.get('seq', -1))
        except Exception:
            return -1

    def _loop(self):
        while not self._stop.is_set():
            try:
                data = json.loads(Path(self.path).read_text(encoding='utf-8'))
                seq  = int(data.get('seq', -1))
                if seq != self._last_seq:
                    self._last_seq = seq
                    self._cb(data)
            except Exception:
                pass
            time.sleep(1.0)


# ──────────────────────────────────────────────────────────────────────────────
# Vibrating-Wire Live Monitor
# ──────────────────────────────────────────────────────────────────────────────

class VWLiveMonitor:
    """
    Live monitor for TWO vibrating-wire load cells (Cell A, Cell B).

    The two cells are independent: each reads its OWN growing LoggerNet TOA5
    .dat file (two separate CR6 loggers, two separate files that may update at
    different rates), keeps its OWN accumulated (timestamp, load) series, and
    draws on its OWN axis. Nothing is shared between cells.

    Every ~2 s (≈0.5 Hz, by design) _tick re-reads each enabled cell's file,
    appends only genuinely new rows, redraws each cell's graph as a single
    time-ordered de-duplicated line, updates the readouts, and logs one row per
    tick to a CSV in a VW_Logs folder next to the watched .dat file.
    """

    POLL_MS = 2000  # ~0.5 Hz update rate (one point per ~2 s) — by design.

    def __init__(self, root: tk.Tk):
        self.root = root
        self.win = tk.Toplevel(root)
        self.win.title("Vibrating-Wire Live Monitor")
        self.win.configure(bg=C['bg'])
        self.win.minsize(980, 620)

        # Per-tick CSV log state (lazily opened once a cell has a folder).
        self._log_fh = None
        self._log_writer = None
        self._log_dir = None
        self._log_start = None   # set when the log is created (Start); elapsed origin
        self._session_tag = datetime.now().strftime('%Y%m%d_%H%M%S')
        # Carry-forward of each cell's latest computed (R1, Load) for the CSV.
        self._last_log = {'A': {'r1': None, 'load': None},
                          'B': {'r1': None, 'load': None}}

        self.cells = {'A': self._new_cell('A'), 'B': self._new_cell('B')}

        self._build_ui()

        self.win.protocol("WM_DELETE_WINDOW", self._on_close)
        self._tick()

    # ── Per-cell state ─────────────────────────────────────────────────────────

    def _new_cell(self, name: str) -> dict:
        return {
            'name':       name,
            'path':       None,
            'columns':    [],
            'series':     {},     # timestamp -> load (dedup by timestamp, keep last)
            'last_ts':    None,   # newest timestamp ingested so far
            # Tk variables (widgets created in _build_cell)
            'enable':     tk.BooleanVar(value=False),
            'path_var':   tk.StringVar(value='(no file selected)'),
            'g_vars':     [tk.StringVar(), tk.StringVar(), tk.StringVar()],
            'g_combos':   [None, None, None],
            'input_mode': tk.StringVar(value='Hz'),       # 'Hz' | 'Digits'
            'equation':   tk.StringVar(value='Linear'),   # 'Linear' | 'Polynomial'
            'const_vars': {k: tk.StringVar(value=str(v))
                           for k, v in VW_DEFAULTS.items()},
            'readout':    tk.StringVar(value='—'),
            'status':     tk.StringVar(value='Disabled'),
            # Matplotlib (created in _build_cell)
            'fig':        None,
            'ax':         None,
            'canvas':     None,
        }

    # ── UI ─────────────────────────────────────────────────────────────────────

    def _btn(self, parent, text, command, style='normal', **kw):
        colors = {
            'primary': (C['btn_primary'], 'white',   '#1D4ED8', 'white'),
            'danger':  (C['btn_danger'],  'white',   '#B91C1C', 'white'),
            'normal':  (C['btn_normal'],  C['text'], '#D1D5DB', C['text']),
        }
        bg, fg, abg, afg = colors.get(style, colors['normal'])
        return tk.Button(parent, text=text, command=command,
                         bg=bg, fg=fg, activebackground=abg, activeforeground=afg,
                         font=('Segoe UI', 9), relief='flat', bd=0,
                         cursor='hand2', padx=8, pady=3, **kw)

    def _build_ui(self):
        header = tk.Frame(self.win, bg=C['accent'], height=34)
        header.pack(fill=tk.X)
        header.pack_propagate(False)
        tk.Label(header, text="Vibrating-Wire Live Monitor — two independent load cells",
                 bg=C['accent'], fg='white',
                 font=('Segoe UI', 11, 'bold')).pack(side=tk.LEFT, padx=12)

        body = tk.Frame(self.win, bg=C['bg'])
        body.pack(fill=tk.BOTH, expand=True)
        body.grid_columnconfigure(0, weight=1, uniform='cell')
        body.grid_columnconfigure(1, weight=1, uniform='cell')
        body.grid_rowconfigure(0, weight=1)

        self._build_cell(body, 'A', column=0)
        self._build_cell(body, 'B', column=1)

    def _build_cell(self, parent, name: str, column: int):
        cell = self.cells[name]

        outer = tk.Frame(parent, bg=C['panel'],
                         highlightbackground=C['border'], highlightthickness=1)
        outer.grid(row=0, column=column, sticky='nsew', padx=4, pady=4)

        # Title + enable
        title = tk.Frame(outer, bg=C['panel'])
        title.pack(fill=tk.X, padx=8, pady=(6, 2))
        tk.Label(title, text=f"Cell {name}", bg=C['panel'], fg=C['accent'],
                 font=('Segoe UI', 12, 'bold')).pack(side=tk.LEFT)
        tk.Checkbutton(title, text="Enable", variable=cell['enable'],
                       bg=C['panel'], fg=C['text'], activebackground=C['panel'],
                       font=('Segoe UI', 9), command=lambda: self._on_enable(name)
                       ).pack(side=tk.RIGHT)

        # File picker
        filerow = tk.Frame(outer, bg=C['panel'])
        filerow.pack(fill=tk.X, padx=8, pady=2)
        self._btn(filerow, "Browse…", lambda: self._browse(name), 'primary',
                  width=9).pack(side=tk.LEFT)
        tk.Label(filerow, textvariable=cell['path_var'], bg=C['panel'],
                 fg=C['muted'], font=('Segoe UI', 8), anchor='w',
                 wraplength=300, justify='left').pack(side=tk.LEFT, padx=6, fill=tk.X)

        # Gauge column dropdowns
        grow = tk.Frame(outer, bg=C['panel'])
        grow.pack(fill=tk.X, padx=8, pady=2)
        tk.Label(grow, text="Gauge cols:", bg=C['panel'], fg=C['text'],
                 font=('Segoe UI', 9)).pack(side=tk.LEFT)
        for i in range(3):
            cb = ttk.Combobox(grow, textvariable=cell['g_vars'][i],
                              values=[], state='readonly', width=10)
            cb.pack(side=tk.LEFT, padx=2)
            cell['g_combos'][i] = cb

        # Input units + equation toggles
        torow = tk.Frame(outer, bg=C['panel'])
        torow.pack(fill=tk.X, padx=8, pady=2)
        tk.Label(torow, text="Input:", bg=C['panel'], fg=C['text'],
                 font=('Segoe UI', 9)).pack(side=tk.LEFT)
        ttk.Combobox(torow, textvariable=cell['input_mode'],
                     values=['Hz', 'Digits'], state='readonly', width=8
                     ).pack(side=tk.LEFT, padx=(2, 10))
        tk.Label(torow, text="Equation:", bg=C['panel'], fg=C['text'],
                 font=('Segoe UI', 9)).pack(side=tk.LEFT)
        ttk.Combobox(torow, textvariable=cell['equation'],
                     values=['Linear', 'Polynomial'], state='readonly', width=11
                     ).pack(side=tk.LEFT, padx=2)

        # Editable constants
        crow = tk.Frame(outer, bg=C['panel'])
        crow.pack(fill=tk.X, padx=8, pady=2)
        tk.Label(crow, text="Constants:", bg=C['panel'], fg=C['muted'],
                 font=('Segoe UI', 8)).pack(side=tk.LEFT)
        for k in ['R0', 'G', 'A', 'B', 'C']:
            tk.Label(crow, text=k, bg=C['panel'], fg=C['text'],
                     font=('Segoe UI', 8)).pack(side=tk.LEFT, padx=(6, 1))
            tk.Entry(crow, textvariable=cell['const_vars'][k], width=8,
                     font=('Segoe UI', 8)).pack(side=tk.LEFT)

        # Large load readout
        readrow = tk.Frame(outer, bg=C['panel'])
        readrow.pack(fill=tk.X, padx=8, pady=(4, 2))
        tk.Label(readrow, textvariable=cell['readout'], bg=C['panel'],
                 fg=C['accent'], font=('Consolas', 26, 'bold')).pack(side=tk.LEFT)
        tk.Label(readrow, text="lbf", bg=C['panel'], fg=C['muted'],
                 font=('Segoe UI', 11)).pack(side=tk.LEFT, padx=(4, 0), pady=(10, 0))
        tk.Label(readrow, textvariable=cell['status'], bg=C['panel'],
                 fg=C['muted'], font=('Segoe UI', 8)).pack(side=tk.RIGHT)

        # Per-cell graph (own figure + own axis)
        fig = Figure(figsize=(4.2, 3.0), dpi=100)
        fig.patch.set_facecolor('#FFFFFF')
        ax = fig.add_subplot(111)
        ax.set_facecolor('#FAFAFA')
        canvas = FigureCanvasTkAgg(fig, master=outer)
        canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True, padx=6, pady=(2, 6))
        cell['fig'], cell['ax'], cell['canvas'] = fig, ax, canvas
        self._draw_cell(name)

    # ── Control callbacks ──────────────────────────────────────────────────────

    def _browse(self, name: str):
        path = filedialog.askopenfilename(
            parent=self.win,
            title=f"Select Cell {name} TOA5 .dat file",
            filetypes=[("TOA5 data", "*.dat"), ("All files", "*.*")])
        if not path:
            return
        cell = self.cells[name]
        cell['path'] = path
        cell['path_var'].set(Path(path).name)
        # Reset this cell's accumulated series — new file, fresh history.
        cell['series'] = {}
        cell['last_ts'] = None
        self._refresh_columns(name)

    def _refresh_columns(self, name: str):
        """Read the file once to populate the gauge-column dropdowns."""
        cell = self.cells[name]
        try:
            df = read_data_snapshot(cell['path'])
        except Exception as e:
            messagebox.showwarning("Read Error",
                                   f"Could not read Cell {name} file:\n{e}",
                                   parent=self.win)
            return
        cols = [c for c in df.columns if c != 'TIMESTAMP']
        cell['columns'] = cols
        for i, cb in enumerate(cell['g_combos']):
            cb['values'] = cols
            # Default to Freq_1/2/3 when present, else first available columns.
            default = VW_FREQ_COLS[i] if VW_FREQ_COLS[i] in cols else (
                cols[i] if i < len(cols) else '')
            cell['g_vars'][i].set(default)

    def _on_enable(self, name: str):
        cell = self.cells[name]
        if cell['enable'].get() and not cell['path']:
            cell['enable'].set(False)
            messagebox.showinfo("No file",
                                f"Select a .dat file for Cell {name} first.",
                                parent=self.win)
            return
        cell['status'].set('Enabled' if cell['enable'].get() else 'Disabled')

    def _consts(self, cell: dict) -> dict:
        """Parse the editable constant entries, falling back to defaults."""
        out = {}
        for k, default in VW_DEFAULTS.items():
            try:
                out[k] = float(cell['const_vars'][k].get())
            except (ValueError, TypeError):
                out[k] = default
        return out

    # ── Live tick ──────────────────────────────────────────────────────────────

    def _tick(self):
        if not self.win.winfo_exists():
            return

        for name in ('A', 'B'):
            try:
                self._update_cell(name)
            except Exception as e:
                self.cells[name]['status'].set(f"Error: {e}")

        self._log_row()

        self.win.after(self.POLL_MS, self._tick)

    def _update_cell(self, name: str):
        cell = self.cells[name]
        if not cell['enable'].get() or not cell['path']:
            return

        df = read_data_snapshot(cell['path'])
        if df is None or len(df) == 0 or 'TIMESTAMP' not in df.columns:
            cell['status'].set('No data yet')
            return

        g_cols = [v.get() for v in cell['g_vars']]
        if any((not c) or (c not in df.columns) for c in g_cols):
            cell['status'].set('Pick 3 valid gauge columns')
            return

        input_is_digits = (cell['input_mode'].get() == 'Digits')
        equation = cell['equation'].get()
        R1, Load = vw_compute(df, g_cols, input_is_digits, equation,
                              self._consts(cell))

        ts = df['TIMESTAMP']
        last_ts = cell['last_ts']

        # Append ONLY genuinely newer rows. A file re-read returns rows we have
        # already ingested; comparing against last_ts ensures we never re-add
        # them. Storing into a dict keyed by timestamp de-duplicates and keeps
        # the last value for any repeated timestamp.
        new_count = 0
        latest_r1 = latest_load = None
        for i in range(len(df)):
            t = ts.iloc[i]
            if pd.isna(t):
                continue
            if last_ts is not None and t <= last_ts:
                continue
            load_val = Load.iloc[i]
            if pd.isna(load_val):
                continue
            cell['series'][t] = float(load_val)
            r1_val = R1.iloc[i]
            latest_r1 = float(r1_val) if pd.notna(r1_val) else latest_r1
            latest_load = float(load_val)
            new_count += 1

        if new_count:
            cell['last_ts'] = max(t for t in [cell['last_ts'], ts.max()]
                                  if t is not None)
            self._last_log[name]['r1'] = latest_r1
            self._last_log[name]['load'] = latest_load

        self._draw_cell(name)

        if cell['series']:
            newest_t = max(cell['series'])
            cell['readout'].set(f"{cell['series'][newest_t]:,.1f}")
            cell['status'].set(
                f"{len(cell['series'])} pts · last {newest_t.strftime('%H:%M:%S')}")
        else:
            cell['status'].set('No data yet')

    def _draw_cell(self, name: str):
        """Redraw a single cell as ONE time-ordered, de-duplicated line on its
        OWN axis using its OWN timestamps (Bug 1 + Bug 2 fix)."""
        cell = self.cells[name]
        ax = cell['ax']
        ax.cla()
        ax.set_facecolor('#FAFAFA')
        ax.set_title(f"Cell {name} — Load vs Time", fontsize=10, color=C['text'])
        ax.set_xlabel("Time", fontsize=8)
        ax.set_ylabel("Load (lbf)", fontsize=8)
        ax.tick_params(labelsize=7)
        ax.grid(True, alpha=0.25, color=C['border'])

        if cell['series']:
            # Sort by timestamp ascending; dict already de-duplicated by
            # timestamp (keep-last). Single series, single line.
            times = sorted(cell['series'])
            loads = [cell['series'][t] for t in times]
            ax.plot(times, loads, color=PLOT_PALETTE[0 if name == 'A' else 1],
                    linewidth=1.6)
            ax.xaxis.set_major_formatter(mdates.DateFormatter('%H:%M:%S'))
            for lbl in ax.get_xticklabels():
                lbl.set_rotation(30)
                lbl.set_ha('right')
        else:
            ax.text(0.5, 0.5, "Waiting for data…", ha='center', va='center',
                    transform=ax.transAxes, fontsize=10, color=C['muted'])

        try:
            cell['fig'].tight_layout()
        except Exception:
            pass
        cell['canvas'].draw()

    # ── CSV logging ────────────────────────────────────────────────────────────

    def _ensure_log(self):
        """Open the per-tick CSV log lazily in a VW_Logs folder next to the
        first enabled cell's .dat file."""
        if self._log_writer is not None:
            return True

        base = None
        for name in ('A', 'B'):
            cell = self.cells[name]
            if cell['enable'].get() and cell['path']:
                base = Path(cell['path']).parent
                break
        if base is None:
            return False

        try:
            log_dir = base / 'VW_Logs'
            log_dir.mkdir(parents=True, exist_ok=True)
            path = log_dir / f"vw_live_log_{self._session_tag}.csv"
            new_file = not path.exists()
            self._log_fh = open(path, 'a', newline='', encoding='utf-8')
            self._log_writer = csv.writer(self._log_fh)
            self._log_dir = log_dir
            if new_file:
                self._log_writer.writerow(
                    ['Timestamp', 'Elapsed_s', 'CellA_R1_digits', 'CellA_Load_lbf',
                     'CellB_R1_digits', 'CellB_Load_lbf'])
                self._log_fh.flush()
            return True
        except Exception:
            self._log_writer = None
            return False

    def _log_row(self):
        """Write one row per tick. Each cell contributes its latest computed
        value (carry-forward); a cell that has produced nothing logs blank."""
        if not (self.cells['A']['enable'].get() or self.cells['B']['enable'].get()):
            return
        if not self._ensure_log():
            return

        # Record the start time when the log is created (Start). The first row
        # is written in the same tick, so its Elapsed_s is exactly 0.0.
        now = datetime.now()
        if self._log_start is None:
            self._log_start = now
        elapsed = round((now - self._log_start).total_seconds(), 3)

        def fmt(v):
            return '' if v is None else f"{v:.4f}"

        a, b = self._last_log['A'], self._last_log['B']
        try:
            self._log_writer.writerow([
                now.strftime('%Y-%m-%d %H:%M:%S'),
                elapsed,
                fmt(a['r1']), fmt(a['load']),
                fmt(b['r1']), fmt(b['load']),
            ])
            self._log_fh.flush()
        except Exception:
            pass

    # ── Lifecycle ──────────────────────────────────────────────────────────────

    def _on_close(self):
        try:
            if self._log_fh is not None:
                self._log_fh.close()
        except Exception:
            pass
        self._log_fh = None
        self._log_writer = None
        self.win.destroy()


# ──────────────────────────────────────────────────────────────────────────────
# Column Mapping Dialog
# ──────────────────────────────────────────────────────────────────────────────

class ColumnMapDialog:
    """
    Modal dialog shown when a file's column names cannot be auto-recognized.
    Displays a preview of the loaded data and lets the user assign which
    column is the X axis and which is the Y axis.

    After win.wait_window() returns, check .cancelled.
    If False, .x_col and .y_col contain the user's selections.
    """

    def __init__(self, parent: tk.Tk, columns: list, preview: pd.DataFrame):
        self.cancelled = True
        self.x_col = columns[0] if columns else ''
        self.y_col = (columns[1] if len(columns) > 1 else columns[0]) if columns else ''

        if not columns:
            return

        win = tk.Toplevel(parent)
        win.title("Assign Column Axes")
        win.configure(bg=C['panel'])
        win.grab_set()
        win.resizable(False, False)
        win.geometry("540x390")

        # ── Header ──────────────────────────────────────────────────────────
        tk.Label(win, text="Column Assignment",
                 bg=C['panel'], fg=C['accent'],
                 font=('Segoe UI', 11, 'bold')).pack(pady=(14, 2), padx=16, anchor='w')
        tk.Label(win,
                 text="The column names in this file weren't auto-recognized.\n"
                      "Use the preview below to identify your columns, then assign axes.",
                 bg=C['panel'], fg=C['text'],
                 font=('Segoe UI', 9), justify='left').pack(padx=16, anchor='w')

        # ── Data preview ────────────────────────────────────────────────────
        pf = tk.Frame(win, bg=C['panel'])
        pf.pack(fill=tk.X, padx=16, pady=(8, 4))
        st = scrolledtext.ScrolledText(
            pf, font=('Consolas', 8), height=8,
            bg=C['accent_lt'], fg=C['text'],
            relief='flat', wrap='none', state='normal')
        st.pack(fill=tk.X)

        col_w = 14
        header_line = "  ".join(f"{str(c)[:col_w]:<{col_w}}" for c in preview.columns)
        sep_line    = "-" * len(header_line)
        st.insert(tk.END, header_line + "\n" + sep_line + "\n")
        for _, row in preview.head(6).iterrows():
            st.insert(tk.END,
                      "  ".join(f"{str(v)[:col_w]:<{col_w}}" for v in row) + "\n")
        st.config(state='disabled')

        # ── Axis selectors ───────────────────────────────────────────────────
        form = tk.Frame(win, bg=C['panel'])
        form.pack(fill=tk.X, padx=16, pady=(6, 4))

        x_var  = tk.StringVar(value=self.x_col)
        y_var  = tk.StringVar(value=self.y_col)
        y_opts = columns + [ALL_CHANNELS_LABEL]

        for label_text, var, opts in [
            ("X Axis  (time / horizontal):", x_var, columns),
            ("Y Axis  (value / vertical):",  y_var, y_opts),
        ]:
            row = tk.Frame(form, bg=C['panel'])
            row.pack(fill=tk.X, pady=3)
            tk.Label(row, text=label_text, bg=C['panel'], fg=C['text'],
                     font=('Segoe UI', 9), width=28, anchor='w').pack(side=tk.LEFT)
            ttk.Combobox(row, textvariable=var, values=opts,
                         state='readonly', width=18).pack(side=tk.LEFT)

        # ── Buttons ──────────────────────────────────────────────────────────
        def _ok():
            self.x_col = x_var.get()
            self.y_col = y_var.get()
            self.cancelled = False
            win.destroy()

        br = tk.Frame(win, bg=C['panel'])
        br.pack(padx=16, pady=(4, 14), fill=tk.X)
        tk.Button(br, text="OK", command=_ok,
                  bg=C['btn_primary'], fg='white',
                  activebackground='#1D4ED8', activeforeground='white',
                  font=('Segoe UI', 9), relief='flat', bd=0,
                  cursor='hand2', padx=10, pady=4).pack(side=tk.LEFT, padx=(0, 6))
        tk.Button(br, text="Cancel", command=win.destroy,
                  bg=C['btn_normal'], fg=C['text'],
                  activebackground='#D1D5DB', activeforeground=C['text'],
                  font=('Segoe UI', 9), relief='flat', bd=0,
                  cursor='hand2', padx=10, pady=4).pack(side=tk.LEFT)

        win.wait_window()


# ──────────────────────────────────────────────────────────────────────────────
# Main Application Window
# ──────────────────────────────────────────────────────────────────────────────

class LabDataApp:

    def __init__(self, root: tk.Tk):
        self.root    = root
        self.dm      = DataManager()
        self.watcher = LiveWatcher(self._on_new_files_thread)
        self.trigger = TriggerWatcher(self._on_trigger_thread)

        # Tkinter variables
        self._preset_var   = tk.StringVar(value=PRESETS[0][0])
        self._x_var        = tk.StringVar()
        self._y_var        = tk.StringVar()
        self._chart_type   = tk.StringVar(value='Line')
        self._show_peak    = tk.BooleanVar(value=False)
        self._live_status   = tk.StringVar(value='Live: OFF')
        self._trigger_status = tk.StringVar(value='Trigger: OFF')
        self._config_name   = tk.StringVar()
        self._compare_mode  = tk.BooleanVar(value=False)
        self._x_label_var   = tk.StringVar()
        self._y_label_var   = tk.StringVar()

        # Camera Monitor window state
        self._monitor_win = None
        self._monitor_cam = 1
        self._monitor_img = None   # held to stop tk.PhotoImage being GC'd

        # Vibrating-Wire Live Monitor window (single instance)
        self._vw_monitor = None

        self._configs = self._read_configs()
        self._recent  = self._read_recent()

        self._build_window()
        self._build_menu()
        self._build_layout()
        self._build_status_bar()
        self._rebuild_recent_menu()
        self._tick_clock()
        self.set_status("Ready — use File → Load CSV File(s) to begin.")


    # ── Window & Menu ────────────────────────────────────────────────────────

    def _build_window(self):
        self.root.title(f"{APP_TITLE}  v{VERSION}")
        self.root.configure(bg=C['bg'])
        self.root.minsize(920, 580)
        try:
            self.root.state('zoomed')
        except Exception:
            self.root.geometry('1100x680')

    def _build_menu(self):
        mb = tk.Menu(self.root)
        self.root.config(menu=mb)

        fm = tk.Menu(mb, tearoff=0)
        fm.add_command(label="Load CSV File(s)…",    command=self.cmd_load,
                       accelerator="Ctrl+O")
        fm.add_command(label="New Session (Clear All)", command=self.cmd_new_session)
        fm.add_separator()
        self._recent_menu = tk.Menu(fm, tearoff=0)
        fm.add_cascade(label="Recent Files", menu=self._recent_menu)
        fm.add_separator()
        fm.add_command(label="Export Current Data…",  command=self.cmd_export,
                       accelerator="Ctrl+S")
        fm.add_command(label="Export as Excel (.xlsx)…", command=self.cmd_export_excel)
        fm.add_command(label="Export Graph as PDF…",     command=self.cmd_export_pdf)
        fm.add_separator()
        fm.add_command(label="Backup Full Data Now",  command=self.cmd_backup)
        fm.add_command(label="Open Backup Folder",    command=self.cmd_open_backup_folder)
        fm.add_separator()
        fm.add_command(label="Exit", command=self.root.quit)
        mb.add_cascade(label="File", menu=fm)

        lm = tk.Menu(mb, tearoff=0)
        lm.add_command(label="Start Live Mode…", command=self.cmd_live_start)
        lm.add_command(label="Stop Live Mode",   command=self.cmd_live_stop)
        lm.add_separator()
        lm.add_command(label="Start Camera Trigger", command=self.cmd_trigger_start)
        lm.add_command(label="Stop Camera Trigger",  command=self.cmd_trigger_stop)
        lm.add_command(label="Camera Monitor…",      command=self.cmd_camera_monitor)
        lm.add_separator()
        lm.add_command(label="VW Live Monitor…",     command=self.cmd_vw_monitor)
        mb.add_cascade(label="Live", menu=lm)

        hm = tk.Menu(mb, tearoff=0)
        hm.add_command(label="How to Use", command=self.cmd_help)
        hm.add_command(label="About",      command=self.cmd_about)
        mb.add_cascade(label="Help", menu=hm)

        self.root.bind('<Control-o>', lambda _: self.cmd_load())
        self.root.bind('<Control-s>', lambda _: self.cmd_export())

    # ── Layout ───────────────────────────────────────────────────────────────

    def _build_layout(self):
        pane = tk.PanedWindow(self.root, orient=tk.HORIZONTAL,
                              bg=C['bg'], sashwidth=5, sashrelief='flat')
        pane.pack(fill=tk.BOTH, expand=True, padx=4, pady=(4, 0))

        self._ctrl_frame  = tk.Frame(pane, bg=C['bg'], width=230)
        self._graph_frame = tk.Frame(pane, bg=C['panel'])
        pane.add(self._ctrl_frame,  minsize=210)
        pane.add(self._graph_frame, minsize=450)

        self._build_controls()
        self._build_graph_area()

    # ─── Controls panel helpers ───────────────────────────────────────────────

    def _section(self, title: str) -> tk.Frame:
        """Return the inner frame of a labeled card section."""
        outer = tk.Frame(self._ctrl_frame, bg=C['bg'])
        outer.pack(fill=tk.X, padx=6, pady=4)

        tk.Label(outer, text=title, bg=C['bg'], fg=C['accent'],
                 font=('Segoe UI', 8, 'bold')).pack(anchor='w', padx=2)

        card = tk.Frame(outer, bg=C['panel'],
                        highlightbackground=C['border'], highlightthickness=1)
        card.pack(fill=tk.X)

        inner = tk.Frame(card, bg=C['panel'])
        inner.pack(fill=tk.X, padx=8, pady=6)
        return inner

    def _btn(self, parent, text, command, style='normal', **kw):
        colors = {
            'primary': (C['btn_primary'], 'white',   '#1D4ED8', 'white'),
            'danger':  (C['btn_danger'],  'white',   '#B91C1C', 'white'),
            'normal':  (C['btn_normal'],  C['text'], '#D1D5DB', C['text']),
        }
        bg, fg, abg, afg = colors.get(style, colors['normal'])
        return tk.Button(parent, text=text, command=command,
                         bg=bg, fg=fg, activebackground=abg, activeforeground=afg,
                         font=('Segoe UI', 9), relief='flat', bd=0,
                         cursor='hand2', padx=8, pady=4, **kw)

    def _label(self, parent, text, muted=False):
        return tk.Label(parent, text=text, bg=C['panel'],
                        fg=C['muted'] if muted else C['text'],
                        font=('Segoe UI', 9))

    def _combobox(self, parent, var, values, width=22):
        cb = ttk.Combobox(parent, textvariable=var, values=values,
                          state='readonly', width=width)
        return cb

    # ─── Controls panel ───────────────────────────────────────────────────────

    def _build_controls(self):
        # Scrollable controls area so nothing gets clipped on small screens
        canvas = tk.Canvas(self._ctrl_frame, bg=C['bg'], highlightthickness=0)
        sb     = ttk.Scrollbar(self._ctrl_frame, orient='vertical', command=canvas.yview)
        canvas.configure(yscrollcommand=sb.set)
        sb.pack(side=tk.RIGHT, fill=tk.Y)
        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self._ctrl_inner = tk.Frame(canvas, bg=C['bg'])
        win_id = canvas.create_window((0, 0), window=self._ctrl_inner, anchor='nw')

        def _on_resize(event):
            canvas.itemconfig(win_id, width=event.width)
        canvas.bind('<Configure>', _on_resize)

        def _on_frame_configure(event):
            canvas.configure(scrollregion=canvas.bbox('all'))
        self._ctrl_inner.bind('<Configure>', _on_frame_configure)

        def _on_mousewheel(event):
            canvas.yview_scroll(-1 * (event.delta // 120), 'units')

        def _bind_scroll(widget):
            """Recursively bind mousewheel to every widget in the controls panel."""
            widget.bind('<MouseWheel>', _on_mousewheel)
            for child in widget.winfo_children():
                _bind_scroll(child)

        canvas.bind('<MouseWheel>', _on_mousewheel)

        # Redirect section() to place inside ctrl_inner
        self._ctrl_frame_target = self._ctrl_inner
        self._build_section_file()
        self._build_section_graph()
        self._build_section_configs()
        self._build_section_chop()
        self._build_section_stats()
        self._build_section_export()

        # Bind scroll after all sections are built so every child widget is covered
        _bind_scroll(self._ctrl_inner)

    def _section(self, title: str) -> tk.Frame:
        outer = tk.Frame(self._ctrl_frame_target, bg=C['bg'])
        outer.pack(fill=tk.X, padx=6, pady=4)
        tk.Label(outer, text=title, bg=C['bg'], fg=C['accent'],
                 font=('Segoe UI', 8, 'bold')).pack(anchor='w', padx=2)
        card = tk.Frame(outer, bg=C['panel'],
                        highlightbackground=C['border'], highlightthickness=1)
        card.pack(fill=tk.X)
        inner = tk.Frame(card, bg=C['panel'])
        inner.pack(fill=tk.X, padx=8, pady=6)
        return inner

    def _build_section_file(self):
        s = self._section("FILE")

        self._btn(s, "Load CSV File(s)", self.cmd_load, 'primary').pack(fill=tk.X, pady=(0, 4))

        tk.Checkbutton(s, text="Compare files (overlay on one graph)",
                       variable=self._compare_mode,
                       bg=C['panel'], fg=C['text'],
                       font=('Segoe UI', 9),
                       activebackground=C['panel']).pack(anchor='w', pady=(0, 4))

        # Live mode row
        row = tk.Frame(s, bg=C['panel'])
        row.pack(fill=tk.X)
        self._live_dot = tk.Label(row, text="●", bg=C['panel'],
                                  fg=C['border'], font=('Segoe UI', 10))
        self._live_dot.pack(side=tk.LEFT)
        tk.Label(row, textvariable=self._live_status,
                 bg=C['panel'], fg=C['muted'],
                 font=('Segoe UI', 8)).pack(side=tk.LEFT, padx=3)
        self._btn(row, "Live…", self.cmd_live_toggle, width=7).pack(side=tk.RIGHT)

        self._file_info_lbl = tk.Label(s, text="No file loaded",
                                       bg=C['panel'], fg=C['muted'],
                                       font=('Segoe UI', 8),
                                       wraplength=190, justify='left')
        self._file_info_lbl.pack(anchor='w', pady=(4, 2))

        self._btn(s, "Map Columns…", self.cmd_map_columns).pack(fill=tk.X, pady=(2, 0))

    def _build_section_graph(self):
        s = self._section("GRAPH")

        self._label(s, "Preset:").pack(anchor='w')
        preset_names = [p[0] for p in PRESETS]
        self._preset_cb = self._combobox(s, self._preset_var, preset_names)
        self._preset_cb.pack(fill=tk.X, pady=(0, 6))
        self._preset_cb.bind('<<ComboboxSelected>>', self._on_preset_select)

        self._label(s, "X Axis:").pack(anchor='w')
        self._x_cb = self._combobox(s, self._x_var, [])
        self._x_cb.pack(fill=tk.X, pady=(0, 6))

        self._label(s, "Y Axis:").pack(anchor='w')
        self._y_cb = self._combobox(s, self._y_var, [])
        self._y_cb.pack(fill=tk.X, pady=(0, 6))

        self._label(s, "Chart Type:").pack(anchor='w')
        self._combobox(s, self._chart_type, CHART_TYPES).pack(fill=tk.X, pady=(0, 6))

        tk.Checkbutton(s, text="Mark Peak Load", variable=self._show_peak,
                       bg=C['panel'], fg=C['text'],
                       font=('Segoe UI', 9),
                       activebackground=C['panel']).pack(anchor='w', pady=(0, 6))

        self._label(s, "X Axis Label (optional):").pack(anchor='w')
        tk.Entry(s, textvariable=self._x_label_var,
                 font=('Segoe UI', 9)).pack(fill=tk.X, pady=(0, 6))

        self._label(s, "Y Axis Label (optional):").pack(anchor='w')
        tk.Entry(s, textvariable=self._y_label_var,
                 font=('Segoe UI', 9)).pack(fill=tk.X, pady=(0, 6))

        self._btn(s, "Update Graph", self.cmd_update_graph, 'primary').pack(fill=tk.X)

    def _build_section_configs(self):
        s = self._section("GRAPH CONFIGS")

        self._label(s, "Save current settings as:").pack(anchor='w')
        row = tk.Frame(s, bg=C['panel'])
        row.pack(fill=tk.X, pady=(0, 6))
        self._config_entry = tk.Entry(row, textvariable=self._config_name,
                                      font=('Segoe UI', 9))
        self._config_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 4))
        self._btn(row, "Save", self.cmd_save_config, 'primary', width=5).pack(side=tk.LEFT)

        self._label(s, "Load saved config:").pack(anchor='w')
        row2 = tk.Frame(s, bg=C['panel'])
        row2.pack(fill=tk.X)
        self._saved_config_cb = ttk.Combobox(row2, state='readonly',
                                              font=('Segoe UI', 9), width=14)
        self._saved_config_cb.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 4))
        self._btn(row2, "Apply", self.cmd_load_config,   width=5).pack(side=tk.LEFT, padx=(0, 2))
        self._btn(row2, "Del",   self.cmd_delete_config, 'danger', width=4).pack(side=tk.LEFT)
        self._refresh_config_list()

    def _build_section_chop(self):
        s = self._section("TIME CHOP")

        self._range_lbl = tk.Label(s, text="Range: —",
                                   bg=C['panel'], fg=C['muted'],
                                   font=('Segoe UI', 8))
        self._range_lbl.pack(anchor='w', pady=(0, 4))

        # Camera-trigger row — fills Start/End automatically on motion
        trow = tk.Frame(s, bg=C['panel'])
        trow.pack(fill=tk.X, pady=(0, 4))
        self._trigger_dot = tk.Label(trow, text="●", bg=C['panel'],
                                     fg=C['border'], font=('Segoe UI', 10))
        self._trigger_dot.pack(side=tk.LEFT)
        tk.Label(trow, textvariable=self._trigger_status,
                 bg=C['panel'], fg=C['muted'],
                 font=('Segoe UI', 8)).pack(side=tk.LEFT, padx=3)
        self._btn(trow, "Trigger…", self.cmd_trigger_toggle,
                  width=8).pack(side=tk.RIGHT)

        for label, attr in [("Start (s):", '_chop_start'), ("End (s):", '_chop_end')]:
            row = tk.Frame(s, bg=C['panel'])
            row.pack(fill=tk.X, pady=2)
            tk.Label(row, text=label, bg=C['panel'], fg=C['text'],
                     font=('Segoe UI', 9), width=9, anchor='w').pack(side=tk.LEFT)
            entry = tk.Entry(row, width=9, font=('Segoe UI', 9))
            entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
            setattr(self, attr, entry)

        row = tk.Frame(s, bg=C['panel'])
        row.pack(fill=tk.X, pady=(6, 0))
        self._btn(row, "Chop",  self.cmd_chop,  'primary', width=8).pack(side=tk.LEFT, padx=(0, 4))
        self._btn(row, "Reset", self.cmd_reset, 'danger',  width=8).pack(side=tk.LEFT)

    def _build_section_stats(self):
        s = self._section("STATISTICS")
        self._stats_text = tk.Text(s, height=7, font=('Consolas', 8),
                                   bg=C['accent_lt'], fg=C['text'],
                                   relief='flat', state='disabled', wrap='none')
        self._stats_text.pack(fill=tk.X)

    def _build_section_export(self):
        s = self._section("EXPORT")
        self._btn(s, "Export CSV",   self.cmd_export,       'primary').pack(fill=tk.X, pady=(0, 3))
        self._btn(s, "Export Excel", self.cmd_export_excel, 'primary').pack(fill=tk.X, pady=(0, 3))
        self._btn(s, "Export PDF",   self.cmd_export_pdf,   'primary').pack(fill=tk.X, pady=(0, 6))

        tk.Frame(s, bg=C['border'], height=1).pack(fill=tk.X, pady=(0, 6))
        self._btn(s, "Backup Full Data Now", self.cmd_backup).pack(fill=tk.X)

    # ── Graph area ───────────────────────────────────────────────────────────

    def _build_graph_area(self):
        self.fig = Figure(figsize=(7, 5), dpi=100)
        self.fig.patch.set_facecolor('#FFFFFF')
        self.ax  = self.fig.add_subplot(111)
        self._style_axes()

        self.canvas = FigureCanvasTkAgg(self.fig, master=self._graph_frame)
        self.canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)

        toolbar_frame = tk.Frame(self._graph_frame, bg=C['panel'])
        toolbar_frame.pack(fill=tk.X)
        toolbar = NavigationToolbar2Tk(self.canvas, toolbar_frame)
        toolbar.update()

        self._draw_placeholder()

    def _style_axes(self):
        self.ax.set_facecolor('#FAFAFA')
        self.ax.tick_params(labelsize=9, colors=C['text'])
        self.ax.spines['top'].set_visible(False)
        self.ax.spines['right'].set_visible(False)
        for sp in ['bottom', 'left']:
            self.ax.spines[sp].set_color(C['border'])

    def _draw_placeholder(self):
        self.ax.clear()
        self._style_axes()
        self.ax.text(0.5, 0.5,
                     'Load a CSV file to begin\n(File → Load CSV File(s)  or  Ctrl+O)',
                     ha='center', va='center', transform=self.ax.transAxes,
                     fontsize=12, color=C['muted'])
        self.ax.axis('off')
        self.canvas.draw()

    # ── Status bar ───────────────────────────────────────────────────────────

    def _build_status_bar(self):
        bar = tk.Frame(self.root, bg=C['accent'], height=24)
        bar.pack(fill=tk.X, side=tk.BOTTOM)
        bar.pack_propagate(False)

        self._status_var = tk.StringVar()
        tk.Label(bar, textvariable=self._status_var,
                 bg=C['accent'], fg='white',
                 font=('Segoe UI', 8), anchor='w').pack(side=tk.LEFT, padx=10)

        self._clock_var = tk.StringVar()
        tk.Label(bar, textvariable=self._clock_var,
                 bg=C['accent'], fg='#93C5FD',
                 font=('Segoe UI', 8)).pack(side=tk.RIGHT, padx=10)

    def _tick_clock(self):
        self._clock_var.set(datetime.now().strftime('%Y-%m-%d  %H:%M:%S'))
        self.root.after(1000, self._tick_clock)

    def set_status(self, msg: str):
        self._status_var.set(f"  {msg}")

    # ── Commands ─────────────────────────────────────────────────────────────

    def cmd_load(self):
        paths = filedialog.askopenfilenames(
            title="Select Data File(s)",
            filetypes=[("Data files", "*.csv *.xlsx *.xlsb *.dat"),
                       ("CSV files", "*.csv"), ("Excel files", "*.xlsx *.xlsb"),
                       ("DAT files", "*.dat"), ("All files", "*.*")]
        )
        if paths:
            self._load_paths(list(paths))

    def cmd_map_columns(self):
        """Open the column-mapping dialog so the user can re-assign X/Y axes."""
        if not self.dm.has_data:
            messagebox.showinfo("No Data", "Load a file first.")
            return
        cols = self.dm.compare_columns if self.dm.compare_frames else self.dm.columns
        dlg = ColumnMapDialog(self.root, cols, self.dm.df_current)
        if not dlg.cancelled:
            self._x_var.set(dlg.x_col)
            self._y_var.set(dlg.y_col)
            self.cmd_update_graph()

    def cmd_new_session(self):
        """Clear all loaded data and return to the blank startup state."""
        if self.dm.has_data:
            if not messagebox.askyesno("New Session",
                    "Clear all loaded data and start fresh?\n\n"
                    "(Your backup files on the Desktop are not affected.)"):
                return
        self.cmd_live_stop()
        self.cmd_trigger_stop()
        self.dm = DataManager()
        self._x_cb['values'] = []
        self._y_cb['values'] = []
        self._x_var.set('')
        self._y_var.set('')
        self._x_label_var.set('')
        self._y_label_var.set('')
        self._range_lbl.config(text="Range: —")
        self._chop_start.delete(0, tk.END)
        self._chop_end.delete(0, tk.END)
        self._file_info_lbl.config(text="No file loaded")
        self._update_stats()
        self._draw_placeholder()
        self.set_status("New session — load a CSV file to begin.")


    def _load_paths(self, paths: list):
        try:
            if self._compare_mode.get():
                errors = self.dm.load_separate(paths)
            else:
                errors = self.dm.load(paths)
                self.dm.compare_frames = []  # clear any previous comparison
                self.dm.compare_names  = []

            try:
                tag = Path(paths[0]).stem[:20]
                self.dm.backup(tag=tag)
                status = f"Loaded {len(paths)} file(s) — auto-backup saved to Desktop."
            except Exception:
                status = f"Loaded {len(paths)} file(s). (Desktop backup failed.)"

            if errors:
                messagebox.showwarning("Load Warnings",
                    "Some files had issues:\n" + "\n".join(errors))

            self._add_recent(paths)
            self._after_load(status)

        except Exception as e:
            messagebox.showerror("Load Error", str(e))

    def _after_load(self, status_msg: str = ""):
        """Refresh all UI controls after new data arrives."""
        cols = self.dm.compare_columns if self.dm.compare_frames else self.dm.columns
        col_with_all = cols + [ALL_CHANNELS_LABEL]

        self._x_cb['values'] = cols
        self._y_cb['values'] = col_with_all

        # Pick x/y defaults — never leave the graph blank when any data exists.
        # x: prefer known time columns, then first available column.
        # y: prefer known measurement columns (load, stress, VW cells),
        #    then first column that differs from x, then x itself.
        _TIME_PREF = ['time_s', 'elapsed_s']
        _MEAS_PREF = ['load_lbf', 'cella_load_lbf', 'cellb_load_lbf',
                      'stress_psi', 'ch3', 'ch4',
                      'cella_r1_digits', 'cellb_r1_digits']

        x = next((c for c in _TIME_PREF if c in cols), cols[0] if cols else '')
        y = next((c for c in _MEAS_PREF if c in cols), None)
        if y is None:
            non_x = [c for c in cols if c != x]
            y = non_x[0] if non_x else x

        self._x_var.set(x)
        self._y_var.set(y)

        # If no recognized time or measurement column was found, prompt the
        # user to assign axes so the graph is never silently blank.
        has_known = any(c in cols for c in _TIME_PREF + _MEAS_PREF)
        if cols and not has_known:
            dlg = ColumnMapDialog(self.root, cols, self.dm.df_current)
            if not dlg.cancelled:
                self._x_var.set(dlg.x_col)
                self._y_var.set(dlg.y_col)

        # Update time-range display and chop fields
        t0, t1 = self.dm.time_range
        self._range_lbl.config(text=f"Range: {t0:.2f}s – {t1:.2f}s  ({len(self.dm.df_current)} rows)")
        self._chop_start.delete(0, tk.END); self._chop_start.insert(0, f"{t0:.2f}")
        self._chop_end.delete(0,   tk.END); self._chop_end.insert(0,   f"{t1:.2f}")

        # File info label
        names = [Path(p).name for p in self.dm.source_files]
        text  = "\n".join(names[:4])
        if len(names) > 4:
            text += f"\n+{len(names)-4} more"
        self._file_info_lbl.config(text=text)

        self._update_stats()
        self.cmd_update_graph()
        if status_msg:
            self.set_status(status_msg)

    def cmd_update_graph(self):
        if not self.dm.has_data:
            return

        x_col = self._x_var.get()
        y_col = self._y_var.get()
        chart = self._chart_type.get()
        peak  = self._show_peak.get()

        self.ax.clear()
        self._style_axes()

        # ── Comparison mode: one series per file ──────────────────────────────
        if self._compare_mode.get() and self.dm.compare_frames:
            x_lbl = self._x_label_var.get().strip() or COL_LABELS.get(x_col, x_col)
            y_lbl = self._y_label_var.get().strip() or COL_LABELS.get(y_col, y_col)

            plotted = 0
            for i, (df, name) in enumerate(zip(self.dm.compare_frames,
                                               self.dm.compare_names)):
                if x_col not in df.columns or y_col not in df.columns:
                    continue
                self._plot_series(df, x_col, y_col, chart,
                                  PLOT_PALETTE[i % len(PLOT_PALETTE)], label=name)
                plotted += 1

            if plotted == 0:
                self.set_status(f"No files have both '{x_col}' and '{y_col}' columns.")
                self.canvas.draw()
                return

            self.ax.set_xlabel(x_lbl, fontsize=10)
            self.ax.set_ylabel(y_lbl, fontsize=10)
            title = f"{y_lbl} vs {x_lbl}  — {plotted} files"
            self.ax.set_title(title, fontsize=11, color=C['text'], pad=8)
            self.ax.legend(fontsize=8, framealpha=0.9)
            self.ax.grid(True, alpha=0.25, color=C['border'])
            self.fig.tight_layout()
            self.canvas.draw()
            self.set_status(f"Comparison: {title}")
            return

        # ── Normal mode: single merged dataset ────────────────────────────────
        df = self.dm.df_current

        if not x_col or x_col not in df.columns:
            self.set_status(f"X column '{x_col}' not found — pick a valid column.")
            self.canvas.draw()
            return

        all_ch = (y_col == ALL_CHANNELS_LABEL)

        if all_ch:
            other = [c for c in df.columns
                     if c != x_col and pd.api.types.is_numeric_dtype(df[c])]
            for i, col in enumerate(other):
                self._plot_series(df, x_col, col, chart,
                                  PLOT_PALETTE[i % len(PLOT_PALETTE)], label=col)
            x_lbl = self._x_label_var.get().strip() or COL_LABELS.get(x_col, x_col)
            self.ax.set_xlabel(x_lbl, fontsize=10)
            self.ax.set_ylabel(self._y_label_var.get().strip() or "Value", fontsize=10)
            title = f"All Channels vs {x_lbl}"
        else:
            if not y_col or y_col not in df.columns:
                self.set_status(f"Y column '{y_col}' not found. Try a different preset or axis.")
                self.canvas.draw()
                return
            self._plot_series(df, x_col, y_col, chart,
                              PLOT_PALETTE[0], show_peak=peak)
            x_lbl = self._x_label_var.get().strip() or COL_LABELS.get(x_col, x_col)
            y_lbl = self._y_label_var.get().strip() or COL_LABELS.get(y_col, y_col)
            self.ax.set_xlabel(x_lbl, fontsize=10)
            self.ax.set_ylabel(y_lbl, fontsize=10)
            title = f"{y_lbl} vs {x_lbl}"

        self.ax.set_title(title, fontsize=11, color=C['text'], pad=8)
        handles, _ = self.ax.get_legend_handles_labels()
        if len(handles) > 1:
            self.ax.legend(fontsize=8, framealpha=0.9)
        self.ax.grid(True, alpha=0.25, color=C['border'])
        self.fig.tight_layout()
        self.canvas.draw()
        self.set_status(f"Graph: {title}  ({len(df)} rows)")

    def _plot_series(self, df, x_col, y_col, chart_type, color,
                     label=None, show_peak=False):
        x   = df[x_col].values
        y   = df[y_col].values
        lbl = label or COL_LABELS.get(y_col, y_col)

        if chart_type == 'Line':
            self.ax.plot(x, y, color=color, linewidth=1.6, label=lbl)
        elif chart_type == 'Scatter':
            self.ax.scatter(x, y, color=color, s=10, label=lbl)
        elif chart_type == 'Area':
            self.ax.fill_between(x, y, alpha=0.20, color=color)
            self.ax.plot(x, y, color=color, linewidth=1.4, label=lbl)

        if show_peak:
            valid = ~np.isnan(y)
            if valid.any():
                idx = int(np.argmax(np.abs(y[valid])))
                px, py = x[valid][idx], y[valid][idx]
                self.ax.axvline(px, color='#DC2626', linestyle='--',
                                linewidth=1, alpha=0.7)
                self.ax.annotate(
                    f"Peak: {py:.1f}",
                    xy=(px, py), xytext=(8, 8), textcoords='offset points',
                    fontsize=8, color='#DC2626',
                    arrowprops=dict(arrowstyle='->', color='#DC2626')
                )

    def _on_preset_select(self, _event=None):
        name = self._preset_var.get()
        for p in PRESETS:
            if p[0] != name:
                continue
            _, x, y, *_ = p
            cols = self.dm.columns
            if x and x in cols:
                self._x_var.set(x)
            if y:
                if y == '__all__':
                    self._y_var.set(ALL_CHANNELS_LABEL)
                elif y in cols:
                    self._y_var.set(y)
            break
        if self.dm.has_data:
            self.cmd_update_graph()

    # ── Chop / Reset ─────────────────────────────────────────────────────────

    def cmd_chop(self):
        try:
            t0 = float(self._chop_start.get())
            t1 = float(self._chop_end.get())
        except ValueError:
            messagebox.showerror("Invalid Input",
                "Start and End times must be numbers (seconds).")
            return

        if t0 >= t1:
            messagebox.showerror("Invalid Range",
                "Start time must be less than end time.")
            return

        try:
            n = self.dm.chop(t0, t1)
        except ValueError as e:
            messagebox.showerror("Chop Error", str(e))
            return

        self._range_lbl.config(
            text=f"Chopped: {t0:.2f}s – {t1:.2f}s  ({n} rows)")
        try:
            self._update_stats()
        except Exception:
            pass
        self.cmd_update_graph()
        self.canvas.draw()
        self.set_status(f"Chopped to {t0}s – {t1}s  ({n} rows retained).")

    def cmd_reset(self):
        try:
            self.dm.reset()
            t0, t1 = self.dm.time_range
            self._range_lbl.config(
                text=f"Range: {t0:.2f}s – {t1:.2f}s  ({len(self.dm.df_current)} rows)")
            self._chop_start.delete(0, tk.END); self._chop_start.insert(0, f"{t0:.2f}")
            self._chop_end.delete(0,   tk.END); self._chop_end.insert(0,   f"{t1:.2f}")
            self._update_stats()
            self.cmd_update_graph()
            self.set_status("Reset to full dataset.")
        except ValueError as e:
            messagebox.showerror("Error", str(e))

    # ── Export / Backup ───────────────────────────────────────────────────────

    def cmd_export(self):
        if not self.dm.has_data:
            messagebox.showinfo("No Data", "Load a file first.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".csv",
            filetypes=[("CSV files", "*.csv")],
            title="Export Data As"
        )
        if not path:
            return
        try:
            self.dm.export(path)
            self.set_status(f"Exported: {Path(path).name}")
        except Exception as e:
            messagebox.showerror("Export Error", str(e))

    def cmd_export_excel(self):
        if not self.dm.has_data:
            messagebox.showinfo("No Data", "Load a file first.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            title="Export as Excel"
        )
        if not path:
            return
        try:
            df = self.dm.df_current

            # Friendly column names for the header row
            rename = {c: COL_LABELS.get(c, c) for c in df.columns}
            df_out = df.rename(columns=rename)

            with pd.ExcelWriter(path, engine='openpyxl') as writer:
                df_out.to_excel(writer, index=False, sheet_name='Data')

                # Remove the blank default sheet openpyxl creates, then
                # make Data the active (first-opened) sheet.
                if 'Sheet' in writer.book.sheetnames:
                    del writer.book['Sheet']
                writer.book.active = writer.sheets['Data']

                # Format the header row bold and auto-size columns
                ws = writer.sheets['Data']
                from openpyxl.styles import Font
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    ws.column_dimensions[col[0].column_letter].width = max_len + 4

                # Metadata sheet
                meta = writer.book.create_sheet('Info')
                meta.append(['Field', 'Value'])
                meta.append(['Exported',    datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
                meta.append(['Source files', ', '.join(
                    Path(p).name for p in self.dm.source_files)])
                t0, t1 = self.dm.time_range
                meta.append(['Time range', f"{t0:.2f}s – {t1:.2f}s"])
                meta.append(['Rows', len(df)])
                stats = self.dm.stats()
                if 'load_lbf' in stats:
                    meta.append(['Peak load (lbf)', stats['load_lbf']['max']])
                for cell in meta[1]:
                    cell.font = Font(bold=True)

            self.set_status(f"Excel exported: {Path(path).name}")
        except Exception as e:
            messagebox.showerror("Export Error", str(e))

    def cmd_export_pdf(self):
        if not self.dm.has_data:
            messagebox.showinfo("No Data", "Load a file first.")
            return
        win = tk.Toplevel(self.root)
        win.title("Export Graph as PDF")
        win.geometry("360x180")
        win.configure(bg=C['panel'])
        win.resizable(False, False)
        win.grab_set()
        tk.Label(win, text="Export Graph as PDF",
                 bg=C['panel'], fg=C['accent'],
                 font=('Segoe UI', 11, 'bold')).pack(pady=(14, 4), padx=16, anchor='w')
        tk.Label(win, text="Test name / label (optional):",
                 bg=C['panel'], fg=C['text'],
                 font=('Segoe UI', 9)).pack(padx=16, anchor='w')
        name_var = tk.StringVar()
        tk.Entry(win, textvariable=name_var,
                 font=('Segoe UI', 9)).pack(fill=tk.X, padx=16, pady=(2, 12))
        def _do_export():
            test_name = name_var.get().strip()
            win.destroy()
            path = filedialog.asksaveasfilename(
                defaultextension=".pdf",
                filetypes=[("PDF files", "*.pdf")],
                title="Save Graph as PDF"
            )
            if not path:
                return
            try:
                self._save_pdf(path, test_name)
                self.set_status(f"PDF exported: {Path(path).name}")
            except Exception as e:
                messagebox.showerror("PDF Export Error", str(e))
        row = tk.Frame(win, bg=C['panel'])
        row.pack(padx=16, fill=tk.X)
        self._btn(row, "Export", _do_export,  'primary', width=10).pack(side=tk.LEFT, padx=(0, 6))
        self._btn(row, "Cancel", win.destroy, width=8).pack(side=tk.LEFT)

    def _save_pdf(self, path: str, test_name: str = ""):
        x_col  = self._x_var.get()
        y_col  = self._y_var.get()
        df     = self.dm.df_current
        fig    = Figure(figsize=(10, 7.5))
        fig.patch.set_facecolor('white')
        ax = fig.add_axes([0.1, 0.22, 0.85, 0.68])
        ax.set_facecolor('#FAFAFA')
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        all_ch = (y_col == ALL_CHANNELS_LABEL)
        if all_ch:
            cols = [c for c in df.columns
                    if c != x_col and pd.api.types.is_numeric_dtype(df[c])]
            for i, col in enumerate(cols):
                ax.plot(df[x_col].values, df[col].values,
                        color=PLOT_PALETTE[i % len(PLOT_PALETTE)],
                        linewidth=1.4, label=col)
            ax.legend(fontsize=9)
        elif self.dm.compare_frames and self._compare_mode.get():
            for i, (frame, name) in enumerate(zip(self.dm.compare_frames,
                                                   self.dm.compare_names)):
                if x_col in frame.columns and y_col in frame.columns:
                    ax.plot(frame[x_col].values, frame[y_col].values,
                            color=PLOT_PALETTE[i % len(PLOT_PALETTE)],
                            linewidth=1.4, label=name)
            ax.legend(fontsize=9)
        else:
            if x_col in df.columns and y_col in df.columns:
                ax.plot(df[x_col].values, df[y_col].values,
                        color=PLOT_PALETTE[0], linewidth=1.6)
        _pdf_xlbl = self._x_label_var.get().strip() or COL_LABELS.get(x_col, x_col)
        _pdf_ylbl = self._y_label_var.get().strip() or COL_LABELS.get(y_col, y_col)
        ax.set_xlabel(_pdf_xlbl, fontsize=11)
        ax.set_ylabel(_pdf_ylbl, fontsize=11)
        ax.grid(True, alpha=0.25)
        title = test_name if test_name else f"{_pdf_ylbl} vs {_pdf_xlbl}"
        ax.set_title(title, fontsize=13, pad=10)
        t0, t1  = self.dm.time_range
        stats   = self.dm.stats()
        peak    = stats.get('load_lbf', {}).get('max', 'N/A')
        sources = ', '.join(Path(p).name for p in self.dm.source_files)
        footer  = (
            f"Source: {sources}     "
            f"Time range: {t0:.2f}s – {t1:.2f}s     "
            f"Rows: {len(df)}     "
            f"Peak load: {peak} lbf     "
            f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )
        fig.text(0.5, 0.04, footer, ha='center', fontsize=7.5, color='#555555', wrap=True)
        fig.text(0.5, 0.01, "ISU Civil Engineering Lab Data Analyzer",
                 ha='center', fontsize=7, color='#999999')
        with PdfPages(path) as pdf:
            pdf.savefig(fig, bbox_inches='tight')
        import matplotlib.pyplot as plt
        plt.close(fig)

    def cmd_backup(self):
        if not self.dm.has_data:
            messagebox.showinfo("No Data", "Load a file first.")
            return
        try:
            out = self.dm.backup()
            self.set_status(f"Backup saved: {Path(out).name}")
            messagebox.showinfo("Backup Complete", f"Saved to:\n{out}")
        except Exception as e:
            messagebox.showerror("Backup Error", str(e))

    def cmd_open_backup_folder(self):
        d = self.dm.backup_dir
        if d.exists():
            os.startfile(str(d))
        else:
            messagebox.showinfo("No Backups",
                f"Backup folder doesn't exist yet:\n{d}")

    # ── Live mode ─────────────────────────────────────────────────────────────

    def cmd_live_toggle(self):
        self.cmd_live_stop() if self.watcher.active else self.cmd_live_start()

    def cmd_live_start(self):
        folder = filedialog.askdirectory(
            title="Select folder to watch for incoming CSV files")
        if not folder:
            return
        self.watcher.start(folder)
        self._live_status.set("Live: ON")
        self._live_dot.config(fg=C['success'])
        self.set_status(f"Live mode ON — watching: {folder}")

    def cmd_live_stop(self):
        self.watcher.stop()
        self._live_status.set("Live: OFF")
        self._live_dot.config(fg=C['border'])
        self.set_status("Live mode stopped.")

    def _on_new_files_thread(self, paths):
        # Called from watcher thread — delegate to main thread via after()
        self.root.after(0, self._handle_new_files, paths)

    def _handle_new_files(self, paths):
        try:
            self.dm.append(paths)
            names = ", ".join(Path(p).name for p in paths)
            self._after_load(f"Live: auto-loaded {names}")
        except Exception as e:
            self.set_status(f"Live load error: {e}")

    # ── Camera trigger ────────────────────────────────────────────────────────

    def cmd_trigger_toggle(self):
        self.cmd_trigger_stop() if self.trigger.active else self.cmd_trigger_start()

    def cmd_trigger_start(self):
        """Begin watching trigger_event.json for camera motion events."""
        self.trigger.start(str(TRIGGER_FILE))
        self._trigger_status.set("Trigger: ON")
        self._trigger_dot.config(fg=C['success'])
        if not TRIGGER_FILE.exists():
            self.set_status(
                "Camera trigger watch ON — waiting for camera_trigger.py "
                f"to create {TRIGGER_FILE.name}.")
        else:
            self.set_status("Camera trigger watch ON — motion will fill "
                            "the Start/End fields.")

    def cmd_trigger_stop(self):
        self.trigger.stop()
        self._trigger_status.set("Trigger: OFF")
        self._trigger_dot.config(fg=C['border'])
        self.set_status("Camera trigger watch stopped.")

    def _on_trigger_thread(self, event):
        # Called from the watcher thread — hand off to the main thread.
        self.root.after(0, self._handle_trigger, event)

    def _handle_trigger(self, event):
        """Auto-populate the chop Start/End fields from a motion event.
        Uses the CR6 elapsed-seconds format (2 decimals) the app uses
        everywhere else."""
        kind = event.get('event')
        try:
            t = float(event.get('elapsed_s'))
        except (TypeError, ValueError):
            return
        cam = event.get('camera', '?')

        if kind == 'start':
            self._chop_start.delete(0, tk.END)
            self._chop_start.insert(0, f"{t:.2f}")
            self.set_status(f"Camera {cam} motion START → chop start = {t:.2f}s")
        elif kind == 'stop':
            self._chop_end.delete(0, tk.END)
            self._chop_end.insert(0, f"{t:.2f}")
            self.set_status(f"Camera {cam} motion STOP → chop end = {t:.2f}s")

    # ── Camera monitor window ─────────────────────────────────────────────────

    def cmd_vw_monitor(self):
        """Open (or focus) the vibrating-wire live monitor window."""
        if (self._vw_monitor is not None
                and self._vw_monitor.win.winfo_exists()):
            self._vw_monitor.win.lift()
            return
        self._vw_monitor = VWLiveMonitor(self.root)

    def cmd_camera_monitor(self):
        """Open (or focus) the live camera monitor window."""
        if self._monitor_win is not None and self._monitor_win.winfo_exists():
            self._monitor_win.lift()
            return

        win = tk.Toplevel(self.root)
        self._monitor_win = win
        win.title("Camera Monitor")
        win.configure(bg=C['panel'])
        win.minsize(500, 420)

        # Header row: camera toggle + RECORDING indicator
        head = tk.Frame(win, bg=C['panel'])
        head.pack(fill=tk.X, padx=10, pady=(10, 4))

        self._monitor_cam_lbl = tk.Label(
            head, text=f"Camera {self._monitor_cam}",
            bg=C['panel'], fg=C['accent'], font=('Segoe UI', 11, 'bold'))
        self._monitor_cam_lbl.pack(side=tk.LEFT)

        self._monitor_rec_lbl = tk.Label(
            head, text="● RECORDING", bg=C['panel'], fg=C['btn_danger'],
            font=('Segoe UI', 10, 'bold'))
        # Packed/forgotten dynamically in _refresh_monitor()

        self._btn(head, "Switch Camera", self._monitor_toggle_cam,
                  'primary', width=14).pack(side=tk.RIGHT)

        # Live feed image
        self._monitor_canvas = tk.Label(win, bg='#000000',
                                        text="Waiting for camera feed…",
                                        fg=C['muted'], font=('Segoe UI', 10))
        self._monitor_canvas.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 6))

        self._monitor_info = tk.Label(
            win, text="", bg=C['panel'], fg=C['muted'],
            font=('Segoe UI', 8), anchor='w', justify='left',
            wraplength=520)
        self._monitor_info.pack(fill=tk.X, padx=10, pady=(0, 10))

        def _on_close():
            self._monitor_win = None
            win.destroy()
        win.protocol("WM_DELETE_WINDOW", _on_close)

        self._refresh_monitor()

    def _monitor_toggle_cam(self):
        self._monitor_cam = 2 if self._monitor_cam == 1 else 1
        self._monitor_cam_lbl.config(text=f"Camera {self._monitor_cam}")

    def _refresh_monitor(self):
        """Update the monitor feed + RECORDING indicator ~10x/sec while the
        window is open. Reads the PNG preview and motion status written by
        camera_trigger.py."""
        win = self._monitor_win
        if win is None or not win.winfo_exists():
            return

        cam     = self._monitor_cam
        preview = CAM_PREVIEW_FILES[cam]
        recording = False
        cam_ok    = False
        info      = "camera_trigger.py is not running — no feed available."

        # Read the live status block for the RECORDING indicator
        try:
            data = json.loads(TRIGGER_FILE.read_text(encoding='utf-8'))
            cinfo     = data.get('cameras', {}).get(str(cam), {})
            recording = bool(cinfo.get('motion'))
            cam_ok    = bool(cinfo.get('ok'))
            elapsed   = data.get('status_elapsed_s', 0.0)
            other     = data.get('cameras', {})
            states    = ", ".join(
                f"Cam {k}: {'MOTION' if v.get('motion') else 'still'}"
                f"{'' if v.get('ok') else ' (no signal)'}"
                for k, v in sorted(other.items()))
            info = f"t = {elapsed:.2f}s    {states}"
        except Exception:
            pass

        # Load the latest preview image (PNG — tk.PhotoImage handles it natively)
        try:
            if preview.exists():
                img = tk.PhotoImage(file=str(preview))
                self._monitor_img = img            # keep a reference
                self._monitor_canvas.config(image=img, text="")
            elif not cam_ok:
                self._monitor_canvas.config(image='',
                    text=f"No feed for Camera {cam}.\n"
                         "Start camera_trigger.py to begin.")
                self._monitor_img = None
        except Exception:
            pass  # keep showing the last good frame on a transient read error

        # Toggle the red RECORDING label
        if recording:
            self._monitor_rec_lbl.pack(side=tk.LEFT, padx=12)
        else:
            self._monitor_rec_lbl.pack_forget()

        self._monitor_info.config(text=info)
        win.after(100, self._refresh_monitor)

    # ── Statistics panel ──────────────────────────────────────────────────────

    def _update_stats(self):
        stats = self.dm.stats()
        lines = []
        for col, s in stats.items():
            lbl = COL_LABELS.get(col, col)
            lines.append(f"{lbl}")
            lines.append(f"  min  {s['min']:>12.3f}")
            lines.append(f"  max  {s['max']:>12.3f}")
            lines.append(f"  mean {s['mean']:>12.3f}")
            lines.append("")
        text = "\n".join(lines).rstrip() or "No data loaded."

        self._stats_text.config(state='normal')
        self._stats_text.delete('1.0', tk.END)
        self._stats_text.insert(tk.END, text)
        self._stats_text.config(state='disabled')

    # ── Graph configs ─────────────────────────────────────────────────────────

    # ── Recent files ──────────────────────────────────────────────────────────

    def _read_recent(self) -> list:
        """Load recent file list from disk. Each entry is a list of paths."""
        try:
            return json.loads(RECENT_FILE.read_text(encoding='utf-8'))
        except Exception:
            return []

    def _write_recent(self):
        RECENT_FILE.write_text(json.dumps(self._recent, indent=2), encoding='utf-8')

    def _add_recent(self, paths: list):
        """Prepend paths to the recent list, remove duplicates, cap at RECENT_MAX."""
        entry = [str(p) for p in paths]
        # Remove any existing entry with the same files
        self._recent = [r for r in self._recent if r != entry]
        self._recent.insert(0, entry)
        self._recent = self._recent[:RECENT_MAX]
        self._write_recent()
        self._rebuild_recent_menu()

    def _rebuild_recent_menu(self):
        """Repopulate the Recent Files submenu from self._recent."""
        self._recent_menu.delete(0, tk.END)
        if not self._recent:
            self._recent_menu.add_command(label="(no recent files)", state='disabled')
            return
        for entry in self._recent:
            # Build a short display label from the filenames
            names = [Path(p).name for p in entry]
            if len(names) == 1:
                label = names[0]
            elif len(names) == 2:
                label = f"{names[0]}, {names[1]}"
            else:
                label = f"{names[0]} + {len(names)-1} more"

            # Use default arg to capture entry value in closure
            self._recent_menu.add_command(
                label=label,
                command=lambda e=entry: self._open_recent(e)
            )
        self._recent_menu.add_separator()
        self._recent_menu.add_command(label="Clear Recent Files",
                                      command=self._clear_recent)

    def _open_recent(self, entry: list):
        """Load a recent entry, removing it from the list if files are missing."""
        missing = [p for p in entry if not Path(p).exists()]
        if missing:
            names = ", ".join(Path(p).name for p in missing)
            if not messagebox.askyesno("Files Not Found",
                    f"These files no longer exist:\n{names}\n\n"
                    "Remove them from Recent Files?"):
                return
            self._recent = [r for r in self._recent if r != entry]
            self._write_recent()
            self._rebuild_recent_menu()
            return
        self._load_paths(entry)

    def _clear_recent(self):
        self._recent = []
        self._write_recent()
        self._rebuild_recent_menu()
        self.set_status("Recent files list cleared.")

    def _read_configs(self) -> dict:
        """Load saved graph configs from disk. Returns empty dict if file missing."""
        try:
            return json.loads(CONFIGS_FILE.read_text(encoding='utf-8'))
        except Exception:
            return {}

    def _write_configs(self):
        """Persist the in-memory configs dict to disk."""
        CONFIGS_FILE.write_text(
            json.dumps(self._configs, indent=2), encoding='utf-8')

    def _refresh_config_list(self):
        """Update the Load dropdown with current saved config names."""
        names = sorted(self._configs.keys())
        self._saved_config_cb['values'] = names
        if names:
            self._saved_config_cb.set(names[0])
        else:
            self._saved_config_cb.set('')

    def cmd_save_config(self):
        name = self._config_name.get().strip()
        if not name:
            messagebox.showwarning("No Name", "Enter a name for this config first.")
            return
        self._configs[name] = {
            'x_col':      self._x_var.get(),
            'y_col':      self._y_var.get(),
            'chart_type': self._chart_type.get(),
            'show_peak':  self._show_peak.get(),
            'x_label':    self._x_label_var.get(),
            'y_label':    self._y_label_var.get(),
        }
        self._write_configs()
        self._refresh_config_list()
        self._saved_config_cb.set(name)
        self.set_status(f"Config saved: \"{name}\"")

    def cmd_load_config(self):
        name = self._saved_config_cb.get()
        if not name or name not in self._configs:
            messagebox.showwarning("No Config", "Select a saved config from the list first.")
            return
        cfg = self._configs[name]
        cols = self.dm.columns + [ALL_CHANNELS_LABEL]

        # Apply saved values — skip any that no longer match loaded columns
        x = cfg.get('x_col', '')
        y = cfg.get('y_col', '')
        if x in self.dm.columns:
            self._x_var.set(x)
        if y in cols:
            self._y_var.set(y)
        self._chart_type.set(cfg.get('chart_type', 'Line'))
        self._show_peak.set(cfg.get('show_peak', False))
        self._x_label_var.set(cfg.get('x_label', ''))
        self._y_label_var.set(cfg.get('y_label', ''))

        if self.dm.has_data:
            self.cmd_update_graph()
        self.set_status(f"Config loaded: \"{name}\"")

    def cmd_delete_config(self):
        name = self._saved_config_cb.get()
        if not name or name not in self._configs:
            return
        if not messagebox.askyesno("Delete Config",
                                   f"Delete saved config \"{name}\"?"):
            return
        del self._configs[name]
        self._write_configs()
        self._refresh_config_list()
        self.set_status(f"Config deleted: \"{name}\"")

    # ── Help / About ─────────────────────────────────────────────────────────

    def cmd_help(self):
        win = tk.Toplevel(self.root)
        win.title("How to Use")
        win.geometry("500x420")
        win.configure(bg=C['panel'])
        win.grab_set()
        win.resizable(False, False)

        tk.Label(win, text="Quick Reference Guide",
                 bg=C['panel'], fg=C['accent'],
                 font=('Segoe UI', 12, 'bold')).pack(pady=(14, 4), padx=16, anchor='w')

        st = scrolledtext.ScrolledText(win, font=('Segoe UI', 9),
                                       bg=C['accent_lt'], fg=C['text'],
                                       relief='flat', wrap='word',
                                       padx=10, pady=8)
        st.pack(fill=tk.BOTH, expand=True, padx=16, pady=4)
        st.insert(tk.END, HELP_TEXT)
        st.config(state='disabled')

        self._btn(win, "Close", win.destroy, 'primary').pack(pady=8)

    def cmd_about(self):
        messagebox.showinfo("About", (
            f"{APP_TITLE}\n"
            f"Version {VERSION}\n\n"
            "Idaho State University\n"
            "Civil Engineering Department\n\n"
            "Processes CSV exports from Campbell Scientific\n"
            "CR6 datalogger via LoggerNet.\n\n"
            "Runs completely offline — no internet required."
        ))


# ──────────────────────────────────────────────────────────────────────────────
# Help text
# ──────────────────────────────────────────────────────────────────────────────

HELP_TEXT = """\
LOADING DATA
  File → Load CSV File(s)  or  Ctrl+O
  Select one or more CSV files. Multiple files are merged and sorted
  by time automatically.
  A backup copy is saved to Desktop\\LabData_Backup\\ on every load.
  Both old format (CH2) and new format (STRESS) columns are recognized.

  File → New Session (Clear All)
  Clears all loaded data and returns to the blank startup state.
  Your backup files on the Desktop are not affected.

GRAPHING
  Choose a Preset from the dropdown for common graph types, or
  manually select X Axis, Y Axis, and Chart Type, then click
  "Update Graph".

  Presets available:
    Load vs Time          — basic load-time trace
    Stress vs Time        — stress vs time (new format files)
    Load vs Displacement  — CH3 as displacement axis
    Stress vs Strain      — CH4 as strain axis
    All Channels vs Time  — all columns on one chart

  Check "Mark Peak Load" to annotate the maximum load point.
  Use the toolbar below the graph to zoom, pan, and save as image.

COMPARING MULTIPLE FILES
  Check "Compare files" in the FILE section before loading.
  Each file is plotted as a separate colored line on the same graph,
  labeled by filename — useful for comparing test runs side by side.
  Uncheck and click Update Graph to return to single-file mode.

GRAPH CONFIGS
  Save your current axis/chart settings under a name using the
  GRAPH CONFIGS section. Load them back anytime with the dropdown.
  Configs are saved to graph_configs.json and survive restarts.

TIME CHOPPING
  Enter Start and End times in seconds, then click Chop.
  Only data within that window will be shown and exported.
  Click Reset to restore the full dataset (backup is unaffected).
  To recover from a bad chop: load the backup from Desktop\\LabData_Backup\\.

LIVE MODE
  Live → Start Live Mode — select the folder where LoggerNet saves
  new CSV files. The app watches that folder and automatically loads
  any new files that appear within 2 seconds.
  A green dot in the FILE section indicates live mode is active.
  Live → Stop Live Mode (or click the Live… button again) to stop.

CAMERA MOTION TRIGGER
  Run camera_trigger.py alongside this app (started at the same moment the
  CR6 test begins, so the elapsed-second clocks line up). It watches two
  PULNiX TM-6701 cameras through the frame grabber and detects motion.

  Live → Start Camera Trigger — the app watches trigger_event.json. When
  motion starts, the Time Chop "Start (s)" field is filled automatically;
  when motion stops, "End (s)" is filled — both in CR6 elapsed seconds.
  A green dot in the TIME CHOP section shows the watch is active.

  Live → Camera Monitor — opens a live feed window. Use "Switch Camera"
  to toggle between Camera 1 and Camera 2. A red RECORDING indicator
  appears whenever motion is active on the viewed camera.

VIBRATING-WIRE LIVE MONITOR
  Live → VW Live Monitor — opens a window for TWO vibrating-wire load
  cells (Cell A, Cell B). The cells are fully independent: each watches
  its OWN growing LoggerNet TOA5 .dat file (two separate CR6 loggers).

  For each cell: Browse to its .dat file, pick the three gauge frequency
  columns (Freq_1/2/3), choose Hz or Digits input, choose the Linear or
  Polynomial equation, edit the calibration constants if needed, then
  tick Enable. The large readout shows the latest load (lbf) and each
  cell draws its own load-vs-time graph from its own file timestamps.

  Updates run about every 2 seconds (≈0.5 Hz) by design. Each tick is
  appended to vw_live_log_<timestamp>.csv inside a VW_Logs folder next
  to the watched .dat file (Timestamp + each cell's R1 digits and load).

EXPORT
  Export CSV   — saves the current working view (post-chop) as CSV.
  Export Excel — saves to .xlsx with bold headers, auto-sized columns,
                 and an Info sheet with metadata (source file, time
                 range, peak load, export date).

BACKUP
  Export PDF   — saves the current graph as a print-ready PDF.
                 You can enter a test name/label that appears as the
                 graph title. Includes a metadata footer.

BACKUP
  Backup saves the FULL original dataset (before any chopping) to
  Desktop\\LabData_Backup\\ with a timestamp. This happens automatically
  on every load, and can also be triggered manually.
"""


# ──────────────────────────────────────────────────────────────────────────────
# Entry point
# ──────────────────────────────────────────────────────────────────────────────

def main():
    root = tk.Tk()
    app  = LabDataApp(root)
    root.protocol("WM_DELETE_WINDOW", root.quit)
    root.mainloop()


if __name__ == '__main__':
    main()
